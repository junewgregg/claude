"""Parse PDF/DOCX/PPTX decks into page-tagged text for evidence citation."""
from __future__ import annotations

import os
from typing import List, Tuple


def parse_pdf(path: str) -> List[Tuple[str, str]]:
    """Return list of (page_label, text). Requires PyMuPDF (fitz)."""
    import fitz  # PyMuPDF

    doc = fitz.open(path)
    out = []
    for i, page in enumerate(doc, start=1):
        text = page.get_text("text").strip()
        if text:
            out.append((f"p.{i}", text))
    doc.close()
    return out


def parse_docx(path: str) -> List[Tuple[str, str]]:
    import docx

    d = docx.Document(path)
    chunks = []
    buf = []
    for p in d.paragraphs:
        if p.text.strip():
            buf.append(p.text.strip())
    for ti, t in enumerate(d.tables, start=1):
        rows = ["| " + " | ".join(c.text.strip() for c in r.cells) + " |" for r in t.rows]
        buf.append(f"[table {ti}]\n" + "\n".join(rows))
    return [("doc", "\n".join(buf))] if buf else []


def parse_pptx(path: str) -> List[Tuple[str, str]]:
    from pptx import Presentation

    prs = Presentation(path)
    out = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
        if texts:
            out.append((f"slide {i}", "\n".join(texts)))
    return out


def parse_deck(path: str) -> List[Tuple[str, str]]:
    """Dispatch on file extension. Returns page-tagged text chunks."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return parse_pdf(path)
    if ext == ".docx":
        return parse_docx(path)
    if ext in (".pptx", ".ppt"):
        return parse_pptx(path)
    if ext in (".txt", ".md"):
        with open(path) as f:
            return [("doc", f.read())]
    raise ValueError(f"Unsupported deck type: {ext}")


def deck_to_prompt_block(path: str, max_chars: int = 60000) -> str:
    """Render a deck as a single page-tagged text block for the model."""
    chunks = parse_deck(path)
    parts = []
    total = 0
    label = os.path.basename(path)
    for tag, text in chunks:
        block = f"[{label} {tag}]\n{text}\n"
        if total + len(block) > max_chars:
            parts.append("[...truncated...]")
            break
        parts.append(block)
        total += len(block)
    return "\n".join(parts)
