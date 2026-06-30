"""Render a TriageScorecard to a .docx report matching the LSB template."""
from __future__ import annotations

from docx import Document
from docx.shared import Pt, RGBColor

from .scorecard import TriageScorecard

_RATING_COLORS = {
    "Acquire": RGBColor(0x1B, 0x7A, 0x3D),
    "Diligence": RGBColor(0xB8, 0x86, 0x0B),
    "Deprioritize": RGBColor(0xA0, 0x30, 0x30),
}


def _section(doc, title):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.rows[0].cells[0]
    run = cell.paragraphs[0].add_run(title)
    run.bold = True
    run.font.size = Pt(12)
    return t


def _kv_table(doc, header, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"
    for j, h in enumerate(header):
        r = t.rows[0].cells[j].paragraphs[0].add_run(str(h))
        r.bold = True
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            t.rows[i].cells[j].text = str(val)
    return t


def render_docx(sc: TriageScorecard, out_path: str) -> str:
    doc = Document()

    # Header
    h = doc.add_paragraph()
    r = h.add_run(f"{sc.asset_name} | Asset evaluation report")
    r.bold = True
    r.font.size = Pt(16)
    doc.add_paragraph(sc.headline_modality).italic = True
    doc.add_paragraph(f"Screening view: {sc.screening_view}")
    doc.add_paragraph(f"Lead asset framing: {sc.lead_asset_framing}")
    doc.add_paragraph(f"Why it lands at {sc.rating.overall_rating.value}: {sc.why_it_lands}")

    # Rating table
    _kv_table(
        doc,
        ["Overall rating", "Disease fit", "Modality fit", "Conviction", "Recommended posture"],
        [[
            sc.rating.overall_rating.value, sc.rating.disease_fit.value,
            sc.rating.modality_fit.value, sc.rating.conviction.value,
            sc.rating.recommended_posture,
        ]],
    )

    _section(doc, "Executive conclusion")
    doc.add_paragraph(sc.executive_conclusion)

    _section(doc, "Asset snapshot")
    s = sc.snapshot
    _kv_table(doc, ["Field", "Value"], [
        ["Asset code", s.asset_code], ["Sponsor / source", s.sponsor_source],
        ["Mechanism", s.mechanism], ["Modality", s.modality],
        ["Lead indication", s.lead_indication], ["Other indications", s.other_indications],
        ["Stage", s.stage], ["Source base", s.source_base],
        ["Priority fit", s.priority_fit], ["Near-term value trigger", s.near_term_value_trigger],
    ])

    _section(doc, "Disease and modality priority fit")
    _kv_table(doc, ["Dimension", "Rating", "What supports the rating", "What still needs to be true"],
              [[p.dimension, p.rating, p.what_supports, p.what_needs_to_be_true] for p in sc.priority_fit])

    _section(doc, "Selected source evidence")
    _kv_table(doc, ["Source page", "Signal", "Decision implication", "Screen read"],
              [[e.source_page, e.signal, e.decision_implication, e.screen_read.value] for e in sc.evidence])

    _section(doc, "Regulatory path, commercial framing, and thesis test")
    _kv_table(doc, ["Thesis gate", "Read", "Commentary", "Implication"],
              [[g.gate, g.read.value, g.commentary, g.implication] for g in sc.thesis_gates])
    doc.add_paragraph(f"Best-fit entry indication: {sc.best_fit_entry_indication}")
    doc.add_paragraph(f"Differentiation claim: {sc.differentiation_claim}")
    doc.add_paragraph(f"Most coherent first story: {sc.most_coherent_first_story}")
    doc.add_paragraph(f"What broadening too early would risk: {sc.broadening_risk}")

    _section(doc, "Key risks, diligence questions, and out-of-rating watchouts")
    doc.add_paragraph("Key risks and unresolved questions").bold = True
    for x in sc.key_risks:
        doc.add_paragraph(x, style="List Bullet")
    doc.add_paragraph("Priority diligence requests").bold = True
    for x in sc.priority_diligence_requests:
        doc.add_paragraph(x, style="List Bullet")
    doc.add_paragraph("Out-of-rating watchouts").bold = True
    for x in sc.out_of_rating_watchouts:
        doc.add_paragraph(x, style="List Bullet")

    _section(doc, "Rating-movement conditions and bottom-line recommendation")
    doc.add_paragraph("Move to Acquire if:").bold = True
    for x in sc.rating_movement.move_to_acquire_if:
        doc.add_paragraph(x, style="List Bullet")
    doc.add_paragraph("Move to Deprioritize if:").bold = True
    for x in sc.rating_movement.move_to_deprioritize_if:
        doc.add_paragraph(x, style="List Bullet")
    bl = doc.add_paragraph()
    br = bl.add_run(f"Bottom line: {sc.bottom_line}")
    br.bold = True

    # Market intelligence section (competitive / SoC / commercial / deal flow)
    if sc.market_intel is not None:
        mi = sc.market_intel
        _section(doc, "Market intelligence: competitive, standard of care, commercial, deal flow")
        _kv_table(doc, ["Lens", "Read"], [
            ["Competitive landscape", mi.competitive_landscape],
            ["Competitive intensity", mi.competitive_intensity.value],
            ["Standard of care", mi.standard_of_care],
            ["Commercial context (CMS)", mi.commercial_context],
            ["Recent deal flow", mi.recent_deal_flow],
        ])
        if mi.exit_acquirer_shortlist:
            doc.add_paragraph("Likely exit acquirers:").bold = True
            for a in mi.exit_acquirer_shortlist:
                doc.add_paragraph(a, style="List Bullet")

    doc.add_paragraph(f"Date prepared: {sc.date_prepared}")

    # Panel determination section
    _section(doc, "Panel determination")
    _kv_table(doc, ["Panel determination", "Key swing factor"],
              [[sc.panel.determination.value, sc.panel.key_swing_factor]])
    doc.add_paragraph(sc.panel.rationale)

    doc.save(out_path)
    return out_path
