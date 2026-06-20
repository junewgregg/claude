"""Local end-to-end tester: email + PDF in, a regular Excel sheet out.

Runs the same brain as the cloud service (extract -> score -> research -> memo)
and the same column layout (including the contact/institute the email carries),
but writes a normal ``.xlsx`` locally and a Markdown memo — no Smartsheet,
SharePoint, or Azure required. Each run appends one asset row to a local state
file and regenerates the sheet, so repeated runs build a comparison table.

Usage:
  # From a raw .eml (body + PDF attachment parsed out of it):
  python -m m365.local_runner --eml sample.eml

  # Or pass the PDF and body directly:
  python -m m365.local_runner --pdf deck.pdf \
      --from "jane@institute.edu" --subject "[Asset] KRAS program" \
      --body "Lead asset attached. Contact: Dr Jane Doe, MIT, 617-555-1212."

  # Add public-database research (slower, more tokens):
  python -m m365.local_runner --eml sample.eml --research
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from rich.console import Console

from biotech_dd import extract as extract_mod
from biotech_dd import research as research_mod
from biotech_dd import score as score_mod
from biotech_dd.report import AssetResult, write_memo
from biotech_dd.schema import Schema, load_schema

from .intake import Attachment, EmailPayload, extract_intake, parse_eml
from .mapping import asset_key, build_row, column_plan

console = Console()

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


# --------------------------------------------------------------------------- #
# Excel output (flat sheet from the shared column plan)
# --------------------------------------------------------------------------- #


def write_local_xlsx(schema: Schema, rows: list[dict], path: str | Path) -> None:
    cols = column_plan(schema)
    score_keys = {c.key: c.scale for c in schema.thesis_scored}

    wb = Workbook()
    ws = wb.active
    ws.title = "DD Table"

    for idx, name in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=idx, value=name)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    score_columns: dict[int, tuple[int, int]] = {}
    for r, row in enumerate(rows, start=2):
        for idx, name in enumerate(cols, start=1):
            cell = ws.cell(row=r, column=idx)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            value = row.get(name, "")
            if name in score_keys:
                try:
                    cell.value = int(value)
                    score_columns[idx] = score_keys[name]
                except (TypeError, ValueError):
                    cell.value = value if value else None
            else:
                cell.value = value or None

    last_row = len(rows) + 1
    for idx, (lo, hi) in score_columns.items():
        col = get_column_letter(idx)
        ws.conditional_formatting.add(
            f"{col}2:{col}{last_row}",
            ColorScaleRule(
                start_type="num", start_value=lo, start_color="F8696B",
                mid_type="num", mid_value=(lo + hi) / 2, mid_color="FFEB84",
                end_type="num", end_value=hi, end_color="63BE7B",
            ),
        )

    for idx in range(1, len(cols) + 1):
        ws.column_dimensions[get_column_letter(idx)].width = 28 if idx > 1 else 22
    ws.freeze_panes = "B2"
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# --------------------------------------------------------------------------- #
# Pipeline (local, no cloud writers)
# --------------------------------------------------------------------------- #


def _field(extracted, key: str) -> str:
    fld = getattr(extracted, key, None) if extracted is not None else None
    return getattr(fld, "value", None) or "" if fld is not None else ""


def run_one(payload: EmailPayload, schema: Schema, thesis: str, out_dir: Path, do_research: bool) -> dict:
    pdfs = payload.pdf_attachments
    if not pdfs:
        raise SystemExit("no PDF attachment found in the email")

    with tempfile.TemporaryDirectory() as tmp:
        deck_path = Path(tmp) / pdfs[0].name
        deck_path.write_bytes(pdfs[0].data)
        console.print("  uploading + extracting deck…")
        deck_file_id = extract_mod.upload_deck(deck_path)
        extracted = extract_mod.extract(schema, deck_file_id)

        console.print("  parsing contact/institute from email…")
        intake = extract_intake(payload.body, payload.sender, deck_file_id)

        console.print("  scoring against thesis…")
        scored = score_mod.score(schema, thesis, extracted)

        researched = None
        if do_research:
            console.print("  researching public databases…")
            researched = research_mod.research(schema, extracted)

    program = _field(extracted, schema.asset_name_field)
    company = intake.company or _field(extracted, "company")
    key = asset_key(program, company, pdfs[0].name.rsplit(".", 1)[0])

    memo_path = out_dir / f"{key}.md"
    write_memo(
        schema,
        AssetResult(
            name=program or key, source_pdf=pdfs[0].name,
            extracted=extracted, scored=scored, researched=researched,
        ),
        memo_path,
    )

    return build_row(
        schema,
        key=key,
        status="New",
        intake=intake,
        extracted=extracted,
        scored=scored,
        researched=researched,
        subject=payload.subject,
        sender=payload.sender,
        intake_date=payload.received or dt.datetime.utcnow().isoformat(),
        thesis_version=hashlib.sha256(thesis.encode()).hexdigest()[:12],
        memo_link=str(memo_path),
        folder_link="",
    )


def _payload_from_args(args) -> EmailPayload:
    if args.eml:
        return parse_eml(args.eml)
    if not args.pdf:
        raise SystemExit("provide --eml, or --pdf with --body/--body-file")
    body = Path(args.body_file).read_text() if args.body_file else args.body
    attachments = [
        Attachment(name=Path(p).name, content_type="application/pdf", data=Path(p).read_bytes())
        for p in args.pdf
    ]
    return EmailPayload(
        subject=args.subject, sender=args.sender, received="", body=body,
        attachments=attachments,
    )


def main(argv: list[str] | None = None) -> None:
    ap = argparse.ArgumentParser(prog="m365.local_runner", description=__doc__)
    ap.add_argument("--eml", help="path to a raw .eml file")
    ap.add_argument("--pdf", action="append", help="PDF attachment path (repeatable)")
    ap.add_argument("--subject", default="[Asset] Local test")
    ap.add_argument("--from", dest="sender", default="tester@example.com")
    ap.add_argument("--body", default="")
    ap.add_argument("--body-file", default="")
    ap.add_argument("--schema", type=Path, default=Path("config/schema.yaml"))
    ap.add_argument("--thesis", type=Path, default=Path("config/thesis.md"))
    ap.add_argument("--out", type=Path, default=Path("out"))
    ap.add_argument("--research", action="store_true", help="run public-database research (slower)")
    args = ap.parse_args(argv)

    schema = load_schema(args.schema)
    thesis = Path(args.thesis).read_text()
    payload = _payload_from_args(args)
    console.print(f"[bold]{payload.subject}[/bold] — {len(payload.pdf_attachments)} PDF(s)")

    row = run_one(payload, schema, thesis, args.out, args.research)

    # Accumulate rows across runs so the sheet grows into a comparison table.
    state_file = args.out / ".local_rows.jsonl"
    state_file.parent.mkdir(parents=True, exist_ok=True)
    existing = {}
    if state_file.exists():
        for line in state_file.read_text().splitlines():
            if line.strip():
                r = json.loads(line)
                existing[r.get("Asset Key", r.get("Source Subject"))] = r
    existing[row.get("Asset Key")] = row  # upsert by asset key
    rows = list(existing.values())
    state_file.write_text("\n".join(json.dumps(r) for r in rows) + "\n")

    xlsx = args.out / "local_table.xlsx"
    write_local_xlsx(schema, rows, xlsx)
    console.print(f"  [green]memo[/green]  -> {args.out / (row['Asset Key'] + '.md')}")
    console.print(f"  [green]sheet[/green] -> {xlsx}  ({len(rows)} asset row(s))")


if __name__ == "__main__":
    main()
