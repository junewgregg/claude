"""Write the DD outputs: an .xlsx comparison table and per-asset Markdown memos.

A single ``AssetResult`` carries one asset's outputs across the three tiers; the
writers below render the collection into a shared spreadsheet (one row per asset)
and an individual memo per asset.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel

from .schema import Schema

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TIER_FILLS = {
    "extracted": PatternFill("solid", fgColor="DDEBF7"),
    "thesis_scored": PatternFill("solid", fgColor="FCE4D6"),
    "researched": PatternFill("solid", fgColor="E2EFDA"),
}


@dataclass
class AssetResult:
    """All outputs for a single asset; later tiers may be ``None`` if skipped."""

    name: str
    source_pdf: str
    extracted: Optional[BaseModel] = None
    scored: Optional[BaseModel] = None
    researched: Optional[BaseModel] = None


def _get(model: Optional[BaseModel], key: str):
    return getattr(model, key, None) if model is not None else None


# --------------------------------------------------------------------------- #
# Spreadsheet
# --------------------------------------------------------------------------- #


def write_xlsx(schema: Schema, results: list[AssetResult], out_path: str | Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "DD Table"

    # Build the column plan: (tier, category) pairs after the leading name column.
    columns: list[tuple[str, object]] = [("_name", "Asset")]
    columns += [("extracted", c) for c in schema.extracted]
    columns += [("thesis_scored", c) for c in schema.thesis_scored]
    columns += [("researched", c) for c in schema.researched]

    # Header row.
    for idx, (tier, cat) in enumerate(columns, start=1):
        header = "Asset" if tier == "_name" else cat.key
        cell = ws.cell(row=1, column=idx, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    # Track score columns so we can color-scale them afterwards.
    score_columns: dict[int, tuple[int, int]] = {}

    for r, result in enumerate(results, start=2):
        for idx, (tier, cat) in enumerate(columns, start=1):
            cell = ws.cell(row=r, column=idx)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

            if tier == "_name":
                cell.value = result.name
                cell.font = Font(bold=True)
                continue

            cell.fill = _TIER_FILLS[tier]
            fld = _get(getattr(result, _attr(tier)), cat.key)

            if tier == "extracted":
                cell.value = _none(getattr(fld, "value", None))
                slide = getattr(fld, "source_slide", None)
                if slide is not None:
                    cell.comment = Comment(f"Slide {slide}", "biotech_dd")
            elif tier == "thesis_scored":
                score = getattr(fld, "score", None)
                label = _enum(getattr(fld, "label", None))
                conf = _enum(getattr(fld, "confidence", None))
                cell.value = score
                summary = " / ".join(p for p in (label, conf) if p)
                rationale = getattr(fld, "rationale", None) or ""
                note = (summary + ("\n\n" if summary and rationale else "") + rationale).strip()
                if note:
                    cell.comment = Comment(note, "biotech_dd")
                score_columns[idx] = cat.scale
            else:  # researched
                cell.value = _none(getattr(fld, "summary", None))
                sources = getattr(fld, "sources", None) or []
                if sources:
                    cell.comment = Comment("\n".join(sources), "biotech_dd")

    # Color-scale each score column over its configured scale.
    last_row = len(results) + 1
    for idx, (lo, hi) in score_columns.items():
        col = get_column_letter(idx)
        rng = f"{col}2:{col}{last_row}"
        ws.conditional_formatting.add(
            rng,
            ColorScaleRule(
                start_type="num", start_value=lo, start_color="F8696B",
                mid_type="num", mid_value=(lo + hi) / 2, mid_color="FFEB84",
                end_type="num", end_value=hi, end_color="63BE7B",
            ),
        )

    _autosize(ws, len(columns))
    ws.freeze_panes = "B2"
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def _attr(tier: str) -> str:
    return {"thesis_scored": "scored"}.get(tier, tier)


def _autosize(ws, ncols: int) -> None:
    for idx in range(1, ncols + 1):
        ws.column_dimensions[get_column_letter(idx)].width = 26 if idx > 1 else 22


# --------------------------------------------------------------------------- #
# Markdown memo
# --------------------------------------------------------------------------- #


def write_memo(schema: Schema, result: AssetResult, out_path: str | Path) -> None:
    lines: list[str] = [f"# Due-Diligence Memo — {result.name}", ""]
    lines.append(f"_Source deck: {result.source_pdf}_")
    lines.append("")

    if result.extracted is not None:
        lines += ["## Extracted facts", "", "| Field | Value | Slide |", "| --- | --- | --- |"]
        for cat in schema.extracted:
            fld = _get(result.extracted, cat.key)
            slide = getattr(fld, "source_slide", None)
            lines.append(
                f"| {cat.key} | {_md(getattr(fld, 'value', None))} | {slide if slide is not None else ''} |"
            )
        lines.append("")

    if result.scored is not None:
        lines += ["## Thesis scorecard", ""]
        for cat in schema.thesis_scored:
            fld = _get(result.scored, cat.key)
            score = getattr(fld, "score", None)
            label = _enum(getattr(fld, "label", None))
            conf = _enum(getattr(fld, "confidence", None))
            header = f"### {cat.key} — {_none(score)}/{cat.scale[1]}"
            tags = " · ".join(p for p in (label, f"confidence: {conf}" if conf else "") if p)
            lines.append(header)
            if tags:
                lines.append(f"*{tags}*")
            lines.append("")
            lines.append(getattr(fld, "rationale", None) or "_No rationale._")
            lines.append("")

    if result.researched is not None:
        lines += ["## Independent research", ""]
        for cat in schema.researched:
            fld = _get(result.researched, cat.key)
            lines.append(f"### {cat.key}")
            lines.append("")
            lines.append(getattr(fld, "summary", None) or "_No findings._")
            sources = getattr(fld, "sources", None) or []
            if sources:
                lines.append("")
                lines.append("**Sources:**")
                lines += [f"- {s}" for s in sources]
            lines.append("")

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    Path(out_path).write_text("\n".join(lines))


# --------------------------------------------------------------------------- #
# Small helpers
# --------------------------------------------------------------------------- #


def _none(value) -> str:
    return "" if value is None else str(value)


def _enum(value) -> str:
    return value.value if value is not None else ""


def _md(value) -> str:
    return _none(value).replace("|", "\\|").replace("\n", " ")
